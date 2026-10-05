"""亞洲大學圖書館個人書單分類號擷取工具。

預設以視窗選擇要填的 Excel，接著開啟 Chrome，由使用者登入後抓取。
--live 可直接處理目前的個人化書單，--input 可在不登入時處理指定檔案。
"""

from __future__ import annotations

import argparse
import csv
from datetime import date
import math
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse
from unicodedata import east_asian_width

import requests
from lxml import html
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_URL = "https://aulib.asia.edu.tw/webpac/"
SHELF_URL = urljoin(BASE_URL, "shelf_personalbook_list.cfm")
LOGIN_URL = urljoin(BASE_URL, "search.cfm")
ROOT = Path(__file__).resolve().parent
INPUT_DIR = ROOT / "待處理"
OUTPUT_DIR = ROOT / "完成"
HEADERS = ("書目資訊", "在架/館藏量", "分類號", "URL")
CALL_NUMBER_PREFIX = re.compile(r"(?<![A-Za-z0-9])([A-Za-z]{0,3}\d+(?:\.\d+)?[A-Za-z0-9]*)")


@dataclass(frozen=True)
class Book:
    description: str
    availability: str
    url: str


def clean_text(value: str) -> str:
    return " ".join((value or "").split())


def valid_book_url(value: str) -> bool:
    parsed = urlparse(value)
    return (
        parsed.scheme == "https"
        and parsed.netloc == "aulib.asia.edu.tw"
        and parsed.path.lower().endswith(("/ele_content.cfm", "/content.cfm"))
        and re.search(r"(?:^|&)mid=\d+(?:&|$)", parsed.query, re.I) is not None
    )


def read_xlsx(path: Path) -> list[Book]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb.active
        header = [clean_text(str(c.value or "")) for c in sheet[1]]
        try:
            title_col = header.index("書目資訊")
            availability_col = header.index("在架/館藏量")
            url_col = header.index("URL")
        except ValueError as exc:
            raise ValueError(f"缺少必要欄位：{exc}") from exc
        books = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            url = str(row[url_col] or "").strip() if len(row) > url_col else ""
            if not url:
                continue
            if not valid_book_url(url):
                raise ValueError(f"無效的書目網址：{url}")
            books.append(Book(
                clean_text(str(row[title_col] or "")),
                clean_text(str(row[availability_col] or "")),
                url,
            ))
        return books
    finally:
        wb.close()


def read_html_xls(path: Path) -> list[Book]:
    source = path.read_text(encoding="utf-8-sig")
    doc = html.fromstring(source)
    rows = doc.xpath("//tr")
    if not rows:
        match = re.search(r'href=["\']([^"\']*sheet\d+\.htm)["\']', source, re.I)
        if not match:
            raise ValueError("此 .xls 是 Excel 網頁封裝，但找不到工作表檔案。")
        sheet_path = (path.parent / match.group(1).replace("/", "\\")).resolve()
        if not sheet_path.is_file():
            raise FileNotFoundError(f"缺少 Excel 附屬資料夾中的工作表：{sheet_path}")
        source = sheet_path.read_text(encoding="utf-8-sig")
        doc = html.fromstring(source)
        rows = doc.xpath("//tr")
    if not rows:
        raise ValueError("工作表沒有資料列。")
    header = [clean_text(c.text_content()) for c in rows[0].xpath("./th|./td")]
    try:
        title_col = header.index("書目資訊")
        availability_col = header.index("在架/館藏量")
        url_col = header.index("URL")
    except ValueError as exc:
        raise ValueError(f"缺少必要欄位：{exc}") from exc
    books = []
    for row in rows[1:]:
        cells = [clean_text(c.text_content()) for c in row.xpath("./th|./td")]
        if len(cells) <= url_col or not cells[url_col]:
            continue
        url = cells[url_col]
        if not valid_book_url(url):
            raise ValueError(f"無效的書目網址：{url}")
        books.append(Book(cells[title_col], cells[availability_col], url))
    return books


def read_input(path: Path) -> list[Book]:
    if path.suffix.lower() == ".xlsx":
        return read_xlsx(path)
    if path.suffix.lower() == ".xls":
        return read_html_xls(path)
    raise ValueError("只支援 .xlsx 或網站匯出的 HTML 格式 .xls。")


def choose_input_gui() -> Path | None:
    """以 Windows 選檔視窗詢問要填入的 Excel；取消時不開始抓取。"""
    import tkinter as tk
    from tkinter import filedialog

    window = tk.Tk()
    window.withdraw()
    window.attributes("-topmost", True)
    try:
        selected = filedialog.askopenfilename(
            parent=window,
            title="選擇要填入分類號的 XLS / XLSX",
            initialdir=str(INPUT_DIR),
            filetypes=[("Excel 檔案", "*.xls *.xlsx"), ("所有檔案", "*.*")],
        )
    finally:
        window.destroy()
    return Path(selected) if selected else None


def extract_holdings(page: bytes) -> tuple[list[str], str]:
    doc = html.fromstring(page)
    found: list[str] = []
    holdings: set[str] = set()
    available = 0
    for table in doc.xpath("//table"):
        head_rows = table.xpath("./thead/tr")
        if not head_rows:
            continue
        labels = [clean_text(c.text_content()) for c in head_rows[0].xpath("./th|./td")]
        if "索書號(卷期)" not in labels:
            continue
        number_col = labels.index("索書號(卷期)")
        accession_col = labels.index("登錄號") if "登錄號" in labels else None
        status_col = labels.index("館藏狀態") if "館藏狀態" in labels else None
        for row_index, row in enumerate(table.xpath("./tbody/tr")):
            cells = row.xpath("./th|./td")
            if number_col >= len(cells):
                continue
            raw = clean_text(cells[number_col].text_content())
            if not raw:
                continue
            match = CALL_NUMBER_PREFIX.search(raw)
            if match and match.group(1) not in found:
                found.append(match.group(1))
            accession = clean_text(cells[accession_col].text_content()) if accession_col is not None and accession_col < len(cells) else ""
            key = accession or f"{id(table)}:{row_index}"
            if key in holdings:
                continue
            holdings.add(key)
            status = clean_text(cells[status_col].text_content()) if status_col is not None and status_col < len(cells) else ""
            if status.startswith("在架"):
                available += 1
    return found, f"{available},{len(holdings)}" if holdings else ""


def book_mid(url: str) -> str:
    return parse_qs(urlparse(url).query).get("mid", [""])[0]


def collect_live_books() -> tuple[list[Book], dict[str, tuple[list[str], str]]]:
    from selenium import webdriver
    from selenium.common.exceptions import NoSuchElementException, NoSuchWindowException, UnexpectedAlertPresentException
    from selenium.webdriver.common.by import By
    from selenium.webdriver.chrome.service import Service
    from selenium.webdriver.support.ui import WebDriverWait
    from webdriver_manager.chrome import ChromeDriverManager

    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()))
    try:
        driver.get(LOGIN_URL)
        login = WebDriverWait(driver, 20).until(
            lambda d: d.find_element(By.ID, "login")
        )
        if "登出" not in login.text:
            trigger = WebDriverWait(driver, 20).until(
                lambda d: d.find_element(By.ID, "login_window_kit_trigger")
            )
            trigger.click()
            print("登入視窗已開啟，請在 Chrome 中自行完成登入；程式會自動繼續。", flush=True)

            def login_complete(browser: webdriver.Chrome) -> bool:
                try:
                    login_area = browser.find_element(By.ID, "login")
                    return (
                        "登出" in login_area.text
                        or (
                            urlparse(browser.current_url).path.endswith("/shelf_personalbook_list.cfm")
                            and bool(browser.find_elements(By.CSS_SELECTOR, "ul.reference-list-content > li"))
                        )
                    )
                except NoSuchElementException:
                    return False
                except UnexpectedAlertPresentException:
                    alert = browser.switch_to.alert
                    print(f"網站提示：{alert.text}", flush=True)
                    alert.accept()
                    return False
                except NoSuchWindowException:
                    handles = browser.window_handles
                    if not handles:
                        raise RuntimeError("登入視窗已關閉，請重新執行程式。")
                    browser.switch_to.window(handles[0])
                    return False

            WebDriverWait(driver, 600, poll_frequency=1).until(
                login_complete
            )
        print("已登入，正在開啟『我的書房 → 個人化書單』。", flush=True)
        driver.get(SHELF_URL)
        WebDriverWait(driver, 20).until(
            lambda d: d.find_elements(By.CSS_SELECTOR, "ul.reference-list-content > li")
            or "登入時間已過" in d.page_source
        )
        if not urlparse(driver.current_url).path.endswith("/shelf_personalbook_list.cfm"):
            raise RuntimeError("未進入『我的書房 → 個人化書單』。請確認已在此 Chrome 視窗登入。")
        summary = clean_text(driver.find_element(By.TAG_NAME, "body").text)
        total_match = re.search(r"第\s*\d+\s*-\s*\d+\s*筆[，,]\s*共\s*(\d+)\s*筆", summary)
        expected_total = int(total_match.group(1)) if total_match else None

        books: list[Book] = []
        clicked: dict[str, tuple[list[str], str]] = {}
        seen: set[str] = set()
        visited_pages: set[tuple[str, ...]] = set()
        for _ in range(100):
            item_rows = driver.find_elements(By.CSS_SELECTOR, "ul.reference-list-content > li")
            signature = tuple(
                box.get_attribute("id")
                for li in item_rows
                for box in li.find_elements(By.CSS_SELECTOR, "input[id^='cart_kit_checkbox_']")
            )
            if signature in visited_pages:
                break
            visited_pages.add(signature)
            page_books: list[Book] = []
            for li in item_rows:
                boxes = li.find_elements(By.CSS_SELECTOR, "input[id^='cart_kit_checkbox_']")
                anchors = li.find_elements(By.CSS_SELECTOR, "a[onclick^='content(']")
                if not boxes or not anchors:
                    continue
                mid = boxes[0].get_attribute("id").removeprefix("cart_kit_checkbox_")
                match = re.search(r"content\('([^']+)'", anchors[0].get_attribute("onclick") or "")
                if not match:
                    continue
                url = urljoin(SHELF_URL, match.group(1))
                if not valid_book_url(url) or url in seen:
                    continue
                if book_mid(url) != mid:
                    raise RuntimeError(f"書目連結與列表 ID 不一致：{mid}")
                seen.add(url)
                title = clean_text(anchors[0].text)
                cells = li.find_elements(By.XPATH, "./div")
                availability = clean_text(cells[-1].text) if cells else ""
                book = Book(title, availability, url)
                books.append(book)
                page_books.append(book)
            for book in page_books:
                mid = book_mid(book.url)
                anchor = driver.find_element(
                    By.XPATH,
                    f"//ul[contains(concat(' ', normalize-space(@class), ' '), ' reference-list-content ')]"
                    f"/li[.//input[@id='cart_kit_checkbox_{mid}']]//a[starts-with(@onclick, 'content(')]",
                )
                anchor.click()
                WebDriverWait(driver, 20).until(
                    lambda d: urlparse(d.current_url).path.endswith("/ele_content.cfm")
                    and book_mid(d.current_url) == mid
                )
                WebDriverWait(driver, 20).until(
                    lambda d: "索書號(卷期)" in d.page_source
                )
                clicked[mid] = extract_holdings(driver.page_source.encode("utf-8"))
                print(f"已點開 [{len(clicked)}] {book.description[:30]}", flush=True)
                driver.back()
                WebDriverWait(driver, 20).until(
                    lambda d: urlparse(d.current_url).path.endswith("/shelf_personalbook_list.cfm")
                    and bool(d.find_elements(By.ID, f"cart_kit_checkbox_{mid}"))
                )
            if expected_total is not None and len(books) >= expected_total:
                break
            next_links = driver.find_elements(By.CSS_SELECTOR, "a[title='下一頁'][onclick]")
            next_link = next((a for a in next_links if a.is_displayed() and a.is_enabled()), None)
            if next_link is None:
                break
            next_link.click()
            try:
                WebDriverWait(driver, 15).until(
                    lambda d: tuple(
                        box.get_attribute("id")
                        for li in d.find_elements(By.CSS_SELECTOR, "ul.reference-list-content > li")
                        for box in li.find_elements(By.CSS_SELECTOR, "input[id^='cart_kit_checkbox_']")
                    ) != signature
                )
            except Exception:
                break

        if not books:
            raise RuntimeError("登入後的個人化書單沒有找到書目。")
        if expected_total is not None and len(books) != expected_total:
            raise RuntimeError(f"個人化書單顯示 {expected_total} 筆，但只擷取 {len(books)} 筆。")
        return books, clicked
    finally:
        driver.quit()


def fetch_results(books: list[Book], session: requests.Session, delay: float) -> tuple[list[list[str]], list[list[str]]]:
    rows: list[list[str]] = []
    issues: list[list[str]] = []
    for position, book in enumerate(books, 1):
        number = ""
        availability = book.availability
        reason = ""
        try:
            response = session.get(book.url, timeout=30)
            if response.status_code >= 400:
                mid = parse_qs(urlparse(book.url).query).get("mid", [""])[0]
                if mid.isdigit():
                    response = session.get(urljoin(BASE_URL, f"ele_content.cfm?mid={mid}"), timeout=30)
            response.raise_for_status()
            found, current_availability = extract_holdings(response.content)
            availability = availability or current_availability
            if not found:
                reason = "未找到索書號(卷期)"
            else:
                number = "、".join(found)
                if len(found) > 1:
                    reason = f"多個不同分類號：{number}"
        except requests.RequestException as exc:
            reason = f"連線失敗：{exc}"
        rows.append([book.description, availability, number, book.url])
        if reason:
            issues.append([str(position), book.description, book.url, reason])
        print(f"[{position}/{len(books)}] {number or '未取得'}  {book.description[:30]}")
        if delay and position < len(books):
            time.sleep(delay)
    return rows, issues


def clicked_results(books: list[Book], clicked: dict[str, tuple[list[str], str]]) -> tuple[list[list[str]], list[list[str]]]:
    missing = [book for book in books if book_mid(book.url) not in clicked]
    if missing:
        raise ValueError(
            f"所選檔案有 {len(missing)} 筆不在目前的個人化書單中；請選擇與網站書單相同的匯出檔。"
        )
    rows: list[list[str]] = []
    issues: list[list[str]] = []
    for position, book in enumerate(books, 1):
        found, current_availability = clicked[book_mid(book.url)]
        number = "、".join(found)
        reason = "未找到索書號(卷期)" if not found else (f"多個不同分類號：{number}" if len(found) > 1 else "")
        rows.append([book.description, book.availability or current_availability, number, book.url])
        if reason:
            issues.append([str(position), book.description, book.url, reason])
    return rows, issues


def save_results(output: Path, rows: list[list[str]], issues: list[list[str]]) -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "工作表1"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
        sheet.cell(sheet.max_row, 4).hyperlink = row[3]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="274A70")
        cell.alignment = Alignment(vertical="center")
    for column, width in {"A": 80, "B": 16, "C": 20, "D": 80}.items():
        sheet.column_dimensions[column].width = width
    for row_number, row in enumerate(sheet.iter_rows(min_row=2), 2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        visual_length = sum(2 if east_asian_width(char) in "WF" else 1 for char in str(row[0].value or ""))
        sheet.row_dimensions[row_number].height = max(29, math.ceil(visual_length / 74) * 18)
        row[2].number_format = "@"  # 分類號是識別碼，保留字母和前導零。
    temporary = output.with_name(output.stem + ".partial.xlsx")
    wb.save(temporary)
    temporary.replace(output)
    issues_path = output.with_name(output.stem + "_待核對.csv")
    if issues:
        with issues_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(("序號", "書目資訊", "URL", "原因"))
            writer.writerows(issues)
    else:
        issues_path.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, help="無登入批次模式：處理指定的 .xls/.xlsx 檔案")
    parser.add_argument("--output", type=Path, help="指定完成檔路徑；可用來更新既有結果")
    parser.add_argument("--live", action="store_true", help="略過選檔視窗，登入後直接處理目前的個人化書單")
    parser.add_argument("--delay", type=float, default=0.5, help="書目頁請求間隔（秒）")
    args = parser.parse_args()
    if args.live and args.input:
        parser.error("--live 與 --input 不可同時使用")
    INPUT_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)
    if args.live:
        books, clicked = collect_live_books()
        today = date.today()
        label = f"個人書單{today.year - 1911:03d}{today.month:02d}{today.day:02d}"
    elif args.input is not None:
        books = read_input(args.input)
        session = requests.Session()
        label = args.input.stem
    else:
        selected = choose_input_gui()
        if selected is None:
            print("未選擇檔案，已取消。")
            return 0
        books = read_input(selected)
        print(f"已選擇：{selected}（{len(books)} 筆）", flush=True)
        _, clicked = collect_live_books()
        label = selected.stem
    if not books:
        raise RuntimeError("來源中沒有書目。")
    if args.input is not None:
        rows, issues = fetch_results(books, session, max(args.delay, 0))
    else:
        rows, issues = clicked_results(books, clicked)
    output = args.output or OUTPUT_DIR / f"{label}_已填分類號.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    save_results(output, rows, issues)
    print(f"完成：{len(rows)} 筆，待核對 {len(issues)} 筆。\n{output}")
    return 0 if not issues else 2


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (OSError, ValueError, RuntimeError) as error:
        print(f"錯誤：{error}", file=sys.stderr)
        sys.exit(1)
